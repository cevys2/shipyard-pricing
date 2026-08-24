"""Parser untuk berkas "REPAIR LIST" / "RINCIAN" -- dokumen kesepakatan di AWAL pekerjaan.

Bedanya dengan `docking_parser.py`: yang itu membaca "REALISASI BIAYA DOCKING", laporan
yang lahir di AKHIR pekerjaan. Repair list lahir di awal dan justru jadi dasar penagihan,
jadi keduanya dokumen yang berbeda dan tidak bisa dibaca parser yang sama. Yang paling
menentukan: di repair list kolom harga satuan berjudul **"SATUAN"**, bukan "HARGA", jadi
`find_harga_group_start(['harga'])` milik docking_parser tidak menemukan apa pun dan
seluruh berkas terbaca nol baris -- gagal tanpa suara, bukan gagal dengan error.

Bentuk berkasnya:

    Baris 1     : "REPAIR LIST" atau "RINCIAN"
    Baris 2     : judul pekerjaan + nama kapal
    Baris 3     : nama klien
    Baris judul : NO | URAIAN PEKERJAAN | VOLUME (VOL, SAT) | SATUAN | TOTAL | KETERANGAN
    Baris angka : 1 | 2 | 3 | 4 | 5 | 6 | 7      <- penomoran kolom, BUKAN data
    Isi         : angka romawi = seksi, huruf A/B/C = sub-seksi, angka 1/2/3 = butir,
                  anak "a." / "1)" / "- " = rincian, kadang berharga kadang tidak
    Penutup     : JUMLAH . PPN 11% . TOTAL

Lima jebakan yang benar-benar ditemui di tiga berkas Basarnas (22 Agustus 2026), semuanya
dijaga oleh `tests/test_repair_list_parser.py`:

1. **Posisi kolom bergeser antar berkas.** Di KN SAR 207 dan ANTAREJA 233 harga ada di
   kolom E; di WIDURA 225 di kolom F, karena grup VOLUME di sana melebar satu kolom.
   Makanya kolom di-scan dari baris judul, tidak pernah ditulis tetap.

2. **Baris penomoran kolom terbaca sebagai data.** Baris `1|2|3|4|5|6|7` tepat di bawah
   baris judul akan masuk sebagai satu baris pekerjaan berharga 5 dan bertotal 6.

3. **Huruf sub-seksi C, D, I, L, M, V, X JUGA angka romawi yang sah.** Ini yang paling
   halus: di KN SAR 207, seksi II punya sub-seksi A, B, C, D. Dibaca apa adanya, "C" dan
   "D" jadi seksi baru, dan seluruh baris Main Deck serta Top Deck kehilangan kategori
   "LAMBUNG KAPAL" -- dan `induk_uraian`-nya ikut kosong. Pemutusnya urutan: sebuah token
   cuma jadi seksi kalau nilai romawinya tepat satu lebih besar dari seksi sebelumnya.
   Setelah seksi II, "C" bernilai 100, bukan 3, jadi dia sub-seksi.

4. **Baris induk sering tidak berharga, anaknya yang berharga.** "2 Pengecatan lambung
   kapal di bawah garis air" kosong, lalu "a. 1 x Cat Primer", "b. 2 x Anti Corrosion",
   "c. 2 x Anti Fouling" masing-masing berharga. Yang diambil yang berharga.

5. **Penulisan penanda anak tidak konsisten.** ANTAREJA menulis `1) Penggantian Part Main
   Engine Tengah`, WIDURA menulis `2  Penggantian Part M/E Kiri` tanpa kurung tutup.

Aturan pengambilan barisnya sendiri sama dengan docking_parser: hanya baris berharga > 0
yang diekstrak. Baris "- ..." tanpa harga adalah rincian isi lump sum induknya, bukan
pekerjaan terpisah -- melewatinya benar, dan yang membuktikannya bukan pendapat melainkan
rekonsiliasi: jumlah baris yang diambil harus persis sama dengan angka JUMLAH di berkas.
"""
import io
import re

import openpyxl

from app.services.docking_parser import (
    PEMISAH_INDUK,
    build_group_labels,
    is_roman,
    norm,
    norm_nospace,
)

# Nilai tiap huruf romawi. Dipakai untuk memutus keambiguan jebakan #3 di atas: yang
# menentukan sebuah token itu seksi atau sub-seksi adalah URUTANNYA, bukan bentuknya.
_NILAI_ROMAWI = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100, "D": 500, "M": 1000}

# Penanda anak di awal teks uraian, dari yang paling dangkal ke yang paling dalam.
# Apostrof di depan ikut dimaafkan: Excel menyimpan sel yang diketik "'- Penggantian
# Terpal" dengan apostrofnya masih menempel, dan itu benar-benar ada di ANTAREJA 233.
_ANAK_HURUF = re.compile(r"^['\"]?\s*([a-z])[.)]\s+(.+)$", re.IGNORECASE)
# Angka harus diikuti tanda baca ATAU minimal dua spasi. Tanpa syarat itu, "2 x Anti
# Fouling (AF)" ikut tertangkap dan uraiannya terpotong jadi "x Anti Fouling (AF)".
_ANAK_ANGKA = re.compile(r"^['\"]?\s*(\d{1,2})(?:[.)]\s*|\s{2,})(.+)$")
_ANAK_STRIP = re.compile(r"^['\"]?\s*[-–•]\s*(.+)$")

# Baris penutup dokumen. Begitu JUMLAH ketemu, isi tabelnya sudah habis.
_PENUTUP = ("jumlah", "ppn", "total", "terbilang")

# Kedalaman baris. Angka romawi (seksi) tidak ikut karena dia jadi kategori, bukan induk.
_LV_SUBSEKSI = 1
_LV_BUTIR = 2
_LV_ANAK_HURUF = 3
_LV_ANAK_ANGKA = 4
_LV_ANAK_STRIP = 5

def nilai_romawi(token: str) -> int:
    """Nilai angka romawi. Token wajib sudah lolos `is_roman()`."""
    t = token.upper()
    total = 0
    for i, huruf in enumerate(t):
        nilai = _NILAI_ROMAWI[huruf]
        berikut = _NILAI_ROMAWI[t[i + 1]] if i + 1 < len(t) else 0
        total += -nilai if berikut > nilai else nilai
    return total


def baca_angka(nilai) -> float | None:
    """Angka dari sel Excel. Mengembalikan None kalau selnya bukan angka.

    Sel harga di berkas ini hampir selalu sudah numerik, tapi tidak selalu: sel yang
    pernah disunting tangan bisa tersimpan sebagai teks "15.000". Titik dibaca sebagai
    pemisah ribuan Indonesia, sejalan dengan `bacaAngkaUang` di frontend.
    """
    if nilai is None:
        return None
    if isinstance(nilai, str):
        teks = nilai.strip()
        if teks in ("", "-"):
            return None
        teks = teks.replace(".", "").replace(",", ".")
        try:
            return float(teks)
        except ValueError:
            return None
    try:
        return float(nilai)
    except (TypeError, ValueError):
        return None


def cari_baris_judul(values) -> int | None:
    """Baris judul tabel: yang memuat "uraian" DAN "total".

    Sengaja tidak memakai kata "harga" seperti docking_parser -- di repair list kolom
    harga satuan justru berjudul "SATUAN", dan kata "harga" tidak muncul sama sekali.
    """
    for i, row in enumerate(values):
        gabung = " ".join(norm_nospace(c) for c in row if c)
        if "uraian" in gabung and "total" in gabung:
            return i
    return None


def _kolom_grup(group_labels: dict[int, str], cocok, sesudah: int | None = None) -> int | None:
    for ci in sorted(group_labels):
        if sesudah is not None and ci <= sesudah:
            continue
        if cocok(group_labels[ci].lower()):
            return ci
    return None


def petakan_kolom(header_row, sub_row) -> dict[str, int | None]:
    """Posisi tiap kolom, dipindai dari baris judul + baris sub-judul.

    Jebakan #1 hidup di sini. Grup VOLUME melebar dua kolom di 207/ANTAREJA dan tiga
    kolom di WIDURA (selnya di-merge), sehingga kolom SATUAN dan TOTAL bergeser ikut.
    Yang dipakai bukan posisi absolut melainkan "grup pertama SESUDAH grup VOLUME".
    """
    group_labels = build_group_labels(header_row)
    sub_labels = {ci: (str(v).strip().lower() if v else "") for ci, v in enumerate(sub_row)}

    uraian = _kolom_grup(group_labels, lambda l: "uraian" in l)
    grup_volume = _kolom_grup(group_labels, lambda l: l.startswith("volume"))

    volume = satuan = None
    akhir_volume = grup_volume
    if grup_volume is not None:
        label_volume = group_labels[grup_volume]
        kolom_volume = [ci for ci, l in group_labels.items() if l == label_volume]
        akhir_volume = max(kolom_volume)
        volume = next((c for c in kolom_volume if sub_labels.get(c, "").startswith("vol")), None)
        satuan = next((c for c in kolom_volume if sub_labels.get(c, "").startswith("sat")), None)

    # "satuan" di sini = HARGA satuan, dan itu memang membingungkan -- tapi begitulah
    # tertulis di ketiga berkas. Dicari sesudah grup VOLUME supaya sub-kolom SAT di
    # dalam grup VOLUME tidak salah terpilih. "harga" tetap diterima kalau suatu saat
    # ada repair list yang menamainya lebih jelas.
    harga = _kolom_grup(
        group_labels, lambda l: l.startswith("satuan") or "harga" in l, sesudah=akhir_volume
    )
    total = _kolom_grup(group_labels, lambda l: "total" in l, sesudah=harga if harga is not None else akhir_volume)
    keterangan = _kolom_grup(
        group_labels, lambda l: l.startswith("ket"), sesudah=total if total is not None else harga
    )
    return {
        "uraian": uraian if uraian is not None else 1,
        "volume": volume,
        "satuan": satuan,
        "harga": harga,
        "total": total,
        "keterangan": keterangan,
    }


def _teks_sel(row, ci: int | None) -> str:
    if ci is None or ci >= len(row) or row[ci] is None:
        return ""
    return str(row[ci]).strip()


def _kedalaman(teks: str) -> tuple[int, str]:
    """Kedalaman baris anak + teksnya tanpa penanda.

    Penandanya dibuang, bukan ikut disimpan: `uraian_pekerjaan` menyimpan apa yang
    dikerjakan, dan "a." cuma nomor urut di dalam satu blok. Konteksnya tidak hilang --
    dia pindah ke `induk_uraian` sebagai teks penuh baris induknya.
    """
    m = _ANAK_HURUF.match(teks)
    if m:
        return _LV_ANAK_HURUF, m.group(2).strip()
    m = _ANAK_ANGKA.match(teks)
    if m:
        return _LV_ANAK_ANGKA, m.group(2).strip()
    m = _ANAK_STRIP.match(teks)
    if m:
        return _LV_ANAK_STRIP, m.group(1).strip()
    # Tanpa penanda sama sekali (mis. "Bongkar pasang shaft propeller" di ANTAREJA).
    # Diperlakukan sebagai anak langsung butir di atasnya.
    return _LV_ANAK_HURUF, teks


def _kepala_dokumen(values, baris_judul: int) -> tuple[str, str, str]:
    """Tiga baris teks di atas tabel: jenis dokumen, judul pekerjaan, nama klien."""
    baris = []
    for row in values[:baris_judul]:
        teks = next((str(c).strip() for c in row if c is not None and str(c).strip()), "")
        if teks:
            baris.append(teks)
    baris += ["", "", ""]
    return baris[0], baris[1], baris[2]


def parse_lembar(values, sheet_name: str) -> dict:
    baris_judul = cari_baris_judul(values)
    if baris_judul is None:
        return {
            "gagal": f"Baris judul tabel (URAIAN ... TOTAL) tidak ketemu di sheet '{sheet_name}'"
        }

    header_row = values[baris_judul]
    sub_row = values[baris_judul + 1] if baris_judul + 1 < len(values) else []
    kolom = petakan_kolom(header_row, sub_row)

    peringatan: list[str] = []
    if kolom["harga"] is None:
        return {"gagal": f"Kolom harga satuan tidak ketemu di sheet '{sheet_name}'"}
    if kolom["volume"] is None:
        peringatan.append("Kolom VOL tidak ketemu -- volume tidak akan terisi.")

    items: list[dict] = []
    kategori: str | None = None
    induk: dict[int, str] = {}
    romawi_terakhir = 0
    jumlah_dok = ppn_dok = total_dok = None

    mulai = baris_judul + 2
    for ri in range(mulai, len(values)):
        row = values[ri]
        if not any(c is not None for c in row):
            continue

        kolom_no = _teks_sel(row, 0)
        teks_uraian = _teks_sel(row, kolom["uraian"])

        harga = baca_angka(_sel(row, kolom["harga"]))
        berharga = harga is not None and harga > 0

        # -- baris penutup: JUMLAH / PPN / TOTAL --------------------------------------
        # Kata kuncinya diterima dari kolom NO. Kalau kolom NO kosong, teks uraian baru
        # ikut diperiksa -- tapi hanya untuk baris yang tidak berharga, supaya pekerjaan
        # sungguhan yang kebetulan diawali kata "Total" tidak ikut terbuang diam-diam.
        kunci = norm(kolom_no) or ("" if berharga else norm(teks_uraian))
        if kunci.startswith(_PENUTUP):
            angka = _angka_penutup(row, kolom["total"])
            if kunci.startswith("jumlah") and jumlah_dok is None:
                jumlah_dok = angka
            elif kunci.startswith("ppn") and ppn_dok is None:
                ppn_dok = angka
            elif kunci.startswith("total") and not teks_uraian and total_dok is None:
                total_dok = angka
            continue

        # -- jebakan #2: baris penomoran kolom ----------------------------------------
        # Cuma diperiksa di baris pertama data. Cirinya sel uraian berisi ANGKA, bukan
        # teks -- baris pekerjaan sungguhan selalu punya uraian berupa teks.
        if ri == mulai and _baris_penomoran(row, kolom["uraian"]):
            continue

        if not teks_uraian:
            continue

        volume = baca_angka(_sel(row, kolom["volume"]))
        satuan = _teks_sel(row, kolom["satuan"]) or None
        total_berkas = baca_angka(_sel(row, kolom["total"]))
        keterangan = _teks_sel(row, kolom["keterangan"])

        def catat(teks: str, rantai: str | None) -> None:
            items.append(
                {
                    "row": ri + 1,
                    "kategori": kategori,
                    "induk_uraian": rantai or None,
                    "uraian": teks,
                    "volume": volume,
                    "satuan": satuan,
                    "harga": harga,
                    "total_berkas": total_berkas,
                    "keterangan": keterangan,
                }
            )

        # -- seksi (angka romawi) -- jebakan #3 ---------------------------------------
        if kolom_no and is_roman(kolom_no):
            nilai = nilai_romawi(kolom_no)
            ambigu = len(kolom_no) == 1 and kolom_no.isalpha()
            if nilai == romawi_terakhir + 1 or not ambigu:
                romawi_terakhir = nilai
                kategori = teks_uraian
                induk = {}
                # Seksi pun bisa berharga: "I DOCKING/ GENERAL SERVICE 1 ls 100.000.000"
                # ada di ketiga berkas dan ikut terhitung di JUMLAH.
                if berharga:
                    catat(teks_uraian, None)
                continue

        # -- sub-seksi (huruf tunggal) -------------------------------------------------
        if kolom_no and len(kolom_no) == 1 and kolom_no.isalpha():
            induk = {_LV_SUBSEKSI: teks_uraian}
            if berharga:
                catat(teks_uraian, None)
            continue

        # -- butir bernomor, atau baris anak -------------------------------------------
        if kolom_no and baca_angka(kolom_no) is not None:
            level, teks_bersih = _LV_BUTIR, teks_uraian
        else:
            level, teks_bersih = _kedalaman(teks_uraian)

        for dalam in [k for k in induk if k >= level]:
            del induk[dalam]
        rantai = PEMISAH_INDUK.join(induk[k] for k in sorted(induk))
        induk[level] = teks_bersih

        if berharga:
            catat(teks_bersih, rantai)

    peringatan += _periksa_baris(items)
    rekonsiliasi = _rekonsiliasi(items, jumlah_dok, ppn_dok, total_dok, peringatan)
    return {
        "sheet_name": sheet_name,
        "items": items,
        "rekonsiliasi": rekonsiliasi,
        "warnings": peringatan,
    }


def _sel(row, ci: int | None):
    return row[ci] if ci is not None and ci < len(row) else None


def _angka_penutup(row, kolom_total: int | None):
    """Angka di baris penutup. Kolom TOTAL dulu; kalau kosong, angka terbesar di baris itu.

    Cadangannya ada karena di sebagian berkas sel JUMLAH di-merge melewati kolom TOTAL,
    sehingga nilainya mendarat di kolom sebelahnya.
    """
    langsung = baca_angka(_sel(row, kolom_total))
    if langsung is not None:
        return langsung
    angka = [n for n in (baca_angka(c) for c in row) if n is not None]
    return max(angka) if angka else None


def _baris_penomoran(row, kolom_uraian: int | None) -> bool:
    sel = _sel(row, kolom_uraian)
    if sel is None or isinstance(sel, str):
        return False
    angka = [n for n in (baca_angka(c) for c in row) if n is not None]
    return len(angka) >= 3 and all(n == int(n) and 1 <= n <= 30 for n in angka)


def _periksa_baris(items: list[dict]) -> list[str]:
    peringatan = []
    for it in items:
        if it["volume"] is None:
            peringatan.append(
                f"Baris {it['row']}: berharga tapi kolom VOL kosong -- nilai barisnya "
                f"dihitung sebagai 1 x harga satuan."
            )
        elif it["total_berkas"] is not None:
            hitung = it["volume"] * it["harga"]
            if abs(hitung - it["total_berkas"]) > 0.5:
                peringatan.append(
                    f"Baris {it['row']}: VOL x harga = {hitung:,.0f} tapi kolom TOTAL di "
                    f"berkas {it['total_berkas']:,.0f} -- selisih {hitung - it['total_berkas']:,.0f}."
                )
    kembar: dict[tuple, list[int]] = {}
    for it in items:
        kembar.setdefault((it["kategori"], it["induk_uraian"], it["uraian"]), []).append(it["row"])
    for (_, _, uraian), baris in kembar.items():
        if len(baris) > 1:
            peringatan.append(
                f"Uraian '{uraian[:60]}' muncul di {len(baris)} baris dengan induk yang sama "
                f"({baris}) -- cek manual, konteksnya belum cukup membedakan."
            )
    return peringatan


def _rekonsiliasi(items, jumlah_dok, ppn_dok, total_dok, peringatan: list[str]) -> dict:
    """Palang verifikasi: jumlah baris yang diambil HARUS sama dengan JUMLAH di dokumen.

    Ini satu-satunya cara aplikasi bisa tahu ada baris yang jatuh waktu impor. Tanpa
    angka volume, palang ini mustahil dipasang -- itulah alasan sebenarnya kolom `volume`
    ditambahkan, bukan sekadar supaya angkanya rapi.

    Yang dijumlahkan VOL x harga satuan, bukan kolom TOTAL di berkas, karena itulah yang
    benar-benar akan tersimpan di katalog. Kalau keduanya berbeda, selisihnya muncul di
    sini -- dan bedanya per baris sudah dilaporkan `_periksa_baris()`.
    """
    terbaca = sum((it["volume"] if it["volume"] is not None else 1) * it["harga"] for it in items)
    selisih = None if jumlah_dok is None else terbaca - jumlah_dok
    if jumlah_dok is None:
        peringatan.append(
            "Baris JUMLAH tidak ketemu di berkas -- impor ini tidak bisa diverifikasi "
            "otomatis terhadap nilai dokumennya."
        )
    elif selisih is not None and abs(selisih) > 0.5:
        peringatan.append(
            f"Jumlah baris terbaca {terbaca:,.0f} TIDAK SAMA dengan JUMLAH di dokumen "
            f"{jumlah_dok:,.0f} (selisih {selisih:,.0f}). Ada baris yang tidak terbaca "
            f"atau terbaca dua kali -- jangan disimpan sebelum ini nol."
        )
    return {
        "jumlah_dokumen": jumlah_dok,
        "ppn_dokumen": ppn_dok,
        "total_dokumen": total_dok,
        "jumlah_terbaca": terbaca,
        "selisih": selisih,
        "cocok": selisih is not None and abs(selisih) <= 0.5,
    }


def _muat_lembar(file_bytes: bytes, filename: str):
    """Semua sheet, apa adanya. Yang dipakai dipilih di `parse_repair_list_file()`."""
    if filename.lower().endswith(".xls"):
        import xlrd

        wb = xlrd.open_workbook(file_contents=file_bytes)
        for nama in wb.sheet_names():
            sheet = wb.sheet_by_name(nama)
            values = [sheet.row_values(r) for r in range(min(sheet.nrows, 5000))]
            yield nama, [[(c if c != "" else None) for c in row] for row in values]
        return
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    for nama in wb.sheetnames:
        ws = wb[nama]
        yield nama, [list(r) for r in ws.iter_rows(min_row=1, max_row=5000, max_col=40, values_only=True)]


def parse_repair_list_file(file_bytes: bytes, filename: str) -> dict:
    """Sheet PERTAMA yang punya baris judul yang bisa dibaca -- bukan yang terakhir.

    docking_parser memakai sheet terakhir karena laporan realisasi memang menaruh
    rekapnya di belakang. Repair list tidak begitu: isinya satu sheet, dan kalau ada
    sheet tambahan biasanya itu lampiran kosong di belakang.

    `detected_tahun` sengaja dikembalikan kosong. Repair list tidak memuat tahun di mana
    pun, dan menebaknya dari nama berkas pernah dicoba di jalur docking lalu terbukti
    salah -- tahun di nama berkas itu tahun terbit dokumen, bukan tahun pekerjaannya.
    Lebih baik kosong dan diisi manusia daripada terisi angka yang keliru.
    """
    gagal: list[str] = []
    for nama, values in _muat_lembar(file_bytes, filename):
        hasil = parse_lembar(values, nama)
        if "gagal" in hasil:
            gagal.append(hasil["gagal"])
            continue
        baris_judul = cari_baris_judul(values)
        jenis, judul, klien = _kepala_dokumen(values, baris_judul)
        hasil.update(
            {
                # Nama sheet, bukan judul dokumen: judulnya kalimat panjang ("PEKERJAAN
                # SPECIAL INSPECTION UNTUK PEMELIHARAAN KAPAL PENYELAMATAN (RESCUE BOAT)
                # KN SAR 207") yang nama kapalnya menempel di ujung tanpa pemisah yang
                # bisa diandalkan. Nama sheet jauh lebih dekat, dan tetap wajib diperiksa
                # manusia di layar pratinjau sebelum disimpan.
                "detected_nama_kapal": nama.strip(),
                "detected_nama_perusahaan": klien,
                "detected_tahun": "",
                "detected_jenis_dokumen": jenis,
                "detected_judul": judul,
            }
        )
        return hasil
    raise ValueError("; ".join(gagal) or "Tidak ada sheet yang bisa dibaca")
