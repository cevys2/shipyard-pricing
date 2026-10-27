"""File Biaya Docking (penawaran pekerjaan induk) dari repair list.

Bentuknya mengikuti Biaya Docking KMP Gerbang Samudra 3 2026 versi final (6 Oktober 2026):
judul, data kapal, description, tabel SCOPE OF WORK persis urutan repair list, lalu tepat di
bawah item terakhir Jumlah / PPN 11% / TOTAL, terbilang, dan tanda tangan.

Kecuali penomoran tabelnya (27 Oktober 2026). GS3 final menaruh semua nomor di kolom A, jadi
"1", "1.1", "a.", lalu "1" lagi di bawah "b." berderet tanpa terlihat mana anak mana. Sekarang
tiap tingkat punya kolom nomornya sendiri (A-E) dan teksnya mulai satu kolom di kanannya,
seperti Biaya Docking Rhama Giri Nusa buatan tangan. Nomornya tetap nomor asli repair list --
dinomori ulang berarti tidak bisa lagi dicocokkan dengan dokumen klien waktu nego.

**Harga satuan sengaja dikosongkan.** Biaya Docking adalah harga PEMBUKA nego, bukan harga
deal -- dan isi katalog justru harga deal. Mengisinya dari katalog berarti menawarkan harga
akhir sebagai harga awal. Harganya diisi di Excel; semua angka di bawahnya ikut lewat rumus.

Karena harga baru diisi SESUDAH file jadi, terbilang tidak bisa ditulis sebagai teks mati --
begitu harga diisi teksnya salah dan tidak ada yang sadar. Excel tidak punya fungsi terbilang
bahasa Indonesia, jadi rumusnya ditaruh di sheet tersembunyi `terbilang`: kolom A berisi kata
untuk 0..999, kolom B memecah TOTAL jadi milyar / juta / ribu / satuan lalu merangkainya.
Halaman cetaknya tetap bersih -- tidak ada lagi blok bantu "Satu, Dua, Tiga" di bawah tabel
seperti di berkas-berkas lama.
"""
import io
import math
import re
from datetime import date, datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Penanda tangan pihak galangan, persis seperti di Biaya Docking GS3 final. Kalau orang atau
# jabatannya berganti, ganti di sini.
PENANDA_TANGAN = ("Alief Febriandy Ramadhan", "VP Business & Marketing")

# A = nomor tingkat 0, B-E = nomor tingkat 1-4, F = teks. B:F selebar B:C lama (62), jadi
# halaman cetaknya tidak ikut melebar.
_LEBAR = {"A": 6, "B": 5, "C": 5, "D": 5, "E": 5, "F": 42, "G": 8, "H": 8, "I": 15, "J": 17, "K": 24}
_TEKS = 6  # F
_TINGKAT_MAKS = _TEKS - 2
_QTY, _SAT, _HARGA, _JUMLAH, _KET = 7, 8, 9, 10, 11
_KOLOM_AKHIR = _KET

_TIPIS = Side(style="thin")
_GANDA = Side(style="double")
_BIRU = PatternFill("solid", fgColor="CAEEFB")  # biru muda kepala tabel, sama dengan GS3
_F_JUDUL = Font(name="Berlin Sans FB Demi", size=16, bold=True)
_F_PITA = Font(name="Berlin Sans FB Demi", size=11, bold=True)
_F_TEKS = Font(name="Calibri", size=11)
_F_TEKS_TEBAL = Font(name="Calibri", size=11, bold=True)
_F_TABEL = Font(name="Arial", size=10)
_F_TABEL_TEBAL = Font(name="Arial", size=10, bold=True)
_AKUNTANSI = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
_TENGAH = Alignment(horizontal="center", vertical="center", wrap_text=True)

_BULAN = [
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
]
_ANGKA = ["", "Satu", "Dua", "Tiga", "Empat", "Lima", "Enam", "Tujuh", "Delapan", "Sembilan",
          "Sepuluh", "Sebelas"]

_KIRI = [
    ("nomor_order", "NOMOR ORDER"),
    ("nama_kapal", "NAMA KAPAL"),
    ("pemilik", "PEMILIK"),
    ("tahun_pembuatan", "TAHUN PEMBUATAN"),
    ("klasifikasi", "KLASIFIKASI KAPAL"),
    ("konstruksi", "KONSTRUKSI KAPAL"),
    ("jenis_kapal", "JENIS/TYPE KAPAL"),
    ("wilayah_operasi", "WILAYAH OPERASI"),
    ("status_survey", "STATUS SURVEY"),
]
# (kunci, nomor, label). Baris tanpa kunci adalah judul kelompok.
_KANAN = [
    (None, "10", "UKURAN UTAMA KAPAL :"),
    ("loa", "", "PANJANG SELURUHNYA ( LOA )"),
    ("lbp", "", "PANJANG GARIS AIR ( LBP )"),
    ("lebar", "", "LEBAR ( B )"),
    ("dalam", "", "DALAM/TINGGI ( D )"),
    ("sarat", "", "DRAFT / Sarat Air ( d )"),
    ("gt", "11", "GT/GRT/NRT"),
    ("daya_mesin", "12", "DAYA MESIN"),
    ("naik_dock", "13", "NAIK DOCK"),
    ("turun_dock", "14", "TURUN DOCK"),
]
_SURAT = ["Nomor", "Tanggal", "Perihal", "Lampiran", "Tanda Tangan", "Jabatan"]


def hari_ini() -> date:
    # Server Railway berjam UTC; tanpa ini file yang dibuat sebelum jam 08.00 WITA bertanggal kemarin.
    return datetime.now(timezone(timedelta(hours=8))).date()


def kata_ratusan(n: int) -> str:
    """0..999 jadi kata: 0 -> "", 115 -> "Seratus Lima Belas"."""
    if n < 12:
        return _ANGKA[n]
    if n < 20:
        return f"{_ANGKA[n - 10]} Belas"
    if n < 100:
        return f"{_ANGKA[n // 10]} Puluh {_ANGKA[n % 10]}".strip()
    depan = "Seratus" if n < 200 else f"{_ANGKA[n // 100]} Ratus"
    return f"{depan} {kata_ratusan(n % 100)}".strip()


def _tanggal_teks(d: date) -> str:
    return f"{d.day:02d} {_BULAN[d.month - 1]} {d.year}"


def _nilai(kapal: dict, kunci: str):
    isi = kapal.get(kunci) or {}
    v = isi.get("nilai")
    if isinstance(v, date):
        v = _tanggal_teks(v)
    elif isinstance(v, float) and v.is_integer():
        v = int(v)
    return v, isi.get("satuan")


def _teks_nilai(kapal: dict, kunci: str) -> str:
    v, satuan = _nilai(kapal, kunci)
    return ":" if v is None else f": {v} {satuan or ''}".rstrip()


def nama_berkas(lingkup: dict, tanggal: date) -> str:
    kapal, _ = _nilai(lingkup["kapal"], "nama_kapal")
    nama = re.sub(r'[\\/:*?"<>|]+', " ", str(kapal or "")).strip()
    return re.sub(r"\s+", " ", f"Biaya Docking {nama} {tanggal.year}") + ".xlsx"


def _tulis(ws, r: int, c: int, nilai=None, font=_F_TEKS, align=None, fmt=None):
    sel = ws.cell(r, c, nilai)
    sel.font = font
    if align:
        sel.alignment = align
    if fmt:
        sel.number_format = fmt
    return sel


def _pita(ws, r: int, teks: str):
    """Baris judul bagian selebar halaman: DATA - DATA KAPAL, DESCRIPTION, SCOPE OF WORK."""
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=_KOLOM_AKHIR)
    _tulis(ws, r, 1, teks, _F_PITA, _TENGAH)
    for c in range(1, _KOLOM_AKHIR + 1):
        ws.cell(r, c).fill = _BIRU
        ws.cell(r, c).border = Border(top=_TIPIS, bottom=_TIPIS)


def _grid(ws, r1: int, r2: int, c1: int = 1, c2: int = _KOLOM_AKHIR):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = Border(left=_TIPIS, right=_TIPIS, top=_TIPIS, bottom=_TIPIS)


def _bingkai(ws, r1: int, r2: int):
    """Bingkai ganda di sekeliling seluruh dokumen; sisi dalam tiap sel dipertahankan."""
    for r in range(r1, r2 + 1):
        for c in range(1, _KOLOM_AKHIR + 1):
            if r not in (r1, r2) and c not in (1, _KOLOM_AKHIR):
                continue
            b = ws.cell(r, c).border
            ws.cell(r, c).border = Border(
                left=_GANDA if c == 1 else b.left,
                right=_GANDA if c == _KOLOM_AKHIR else b.right,
                top=_GANDA if r == r1 else b.top,
                bottom=_GANDA if r == r2 else b.bottom,
            )


def _tinggi(teks: str, kolom_awal: int) -> float:
    # ponytail: perkiraan dari jumlah huruf (~70 huruf Arial 10 per baris selebar B:F = 62,
    # diukur dari PDF Excel). Excel tidak menyesuaikan tinggi sel gabungan sendiri; kalau
    # sesekali kurang, tarik di Excel.
    lebar = sum(_LEBAR[chr(64 + c)] for c in range(kolom_awal, _TEKS + 1))
    baris = max(1, math.ceil(len(teks) / max(20, 70 * lebar / 62)))
    return max(15.0, 13.0 * baris + 2)


def _kepala(ws, kapal: dict, tanggal: date) -> int:
    """Judul s.d. kepala tabel. Mengembalikan nomor baris pertama isi tabel."""
    for r, teks, font in ((1, "BIAYA DOCKING", _F_JUDUL), (2, "( Pekerjaan Induk )", _F_TEKS_TEBAL)):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=_KOLOM_AKHIR)
        _tulis(ws, r, 1, teks, font, _TENGAH)
    ws.row_dimensions[1].height = 24

    _tulis(ws, 4, 1, "No. Revisi :")
    _tulis(ws, 4, _HARGA, "Tanggal Terbit :", align=Alignment(horizontal="right"))
    ws.merge_cells("J4:K4")
    _tulis(ws, 4, _JUMLAH, tanggal, fmt="[$-421]dd mmmm yyyy;@", align=Alignment(horizontal="left"))

    # Label di B:E (selebar kolom B lama), nilai di F.
    _pita(ws, 5, "DATA - DATA KAPAL")
    for k, (kunci, label) in enumerate(_KIRI):
        r = 6 + k
        _tulis(ws, r, 1, k + 1, align=Alignment(horizontal="center"))
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        _tulis(ws, r, 2, label)
        _tulis(ws, r, _TEKS, _teks_nilai(kapal, kunci), _F_TEKS_TEBAL)
    for k, (kunci, nomor, label) in enumerate(_KANAN):
        r = 6 + k
        if nomor:
            _tulis(ws, r, _QTY, int(nomor), align=Alignment(horizontal="center"))
        ws.merge_cells(start_row=r, start_column=_SAT, end_row=r, end_column=_JUMLAH)
        _tulis(ws, r, _SAT, label, _F_TEKS_TEBAL if kunci is None else _F_TEKS)
        if kunci:
            _tulis(ws, r, _KET, _teks_nilai(kapal, kunci), _F_TEKS_TEBAL)

    nama, _ = _nilai(kapal, "nama_kapal")
    pemilik, _ = _nilai(kapal, "pemilik")
    _pita(ws, 17, "DESCRIPTION")
    _tulis(ws, 18, 1, 1, _F_TEKS_TEBAL, Alignment(horizontal="center"))
    ws.merge_cells("B18:K18")
    _tulis(ws, 18, 2, f'Biaya pelaksanaan pekerjaan Docking kapal "{nama or ""}", berdasarkan :', _F_TEKS_TEBAL)
    ws.merge_cells("B19:E19")
    _tulis(ws, 19, 2, "1.1", _F_TEKS_TEBAL, Alignment(horizontal="center"))
    _tulis(ws, 19, _TEKS, f"Surat dari {pemilik or ''}".rstrip(), _F_TEKS_TEBAL)
    for k, label in enumerate(_SURAT):
        ws.merge_cells(start_row=20 + k, start_column=2, end_row=20 + k, end_column=5)
        _tulis(ws, 20 + k, 2, label, align=Alignment(horizontal="right", indent=1))
        _tulis(ws, 20 + k, _TEKS, ":")

    _pita(ws, 27, "SCOPE OF WORK")
    for rentang, teks in (("A28:A29", "NO."), ("B28:F29", "U R A I A N"), ("G28:G29", "QTY"),
                          ("H28:H29", "SAT."), ("I28:J28", "HARGA (Rp.)"), ("K28:K29", "KETERANGAN")):
        ws.merge_cells(rentang)
        _tulis(ws, int(rentang[1:3]), ord(rentang[0]) - 64, teks, _F_PITA, _TENGAH)
    _tulis(ws, 29, _HARGA, "Satuan", _F_PITA, _TENGAH)
    _tulis(ws, 29, _JUMLAH, "Jumlah", _F_PITA, _TENGAH)
    _grid(ws, 28, 29)
    for r in (28, 29):
        for c in range(1, _KOLOM_AKHIR + 1):
            ws.cell(r, c).fill = _BIRU
    return 30


def _isi_tabel(ws, baris: list[dict], r0: int) -> int:
    """Baris pekerjaan persis urutan repair list. Mengembalikan baris terakhirnya."""
    for k, b in enumerate(baris):
        r = r0 + k
        font = _F_TABEL_TEBAL if b["seksi"] else _F_TABEL
        # Nomor di kolom tingkatnya, teks mulai satu kolom di kanannya sampai F.
        # ponytail: lima kolom nomor; tingkat 5 ke atas (belum pernah ada) ditaruh di tingkat 4.
        t = min(b["tingkat"], _TINGKAT_MAKS)
        _tulis(ws, r, 1 + t, b["no"] or None, font,
               Alignment(horizontal="center" if t == 0 else "left", vertical="top"))
        if 2 + t < _TEKS:
            ws.merge_cells(start_row=r, start_column=2 + t, end_row=r, end_column=_TEKS)
        _tulis(ws, r, 2 + t, b["uraian"], font, Alignment(vertical="top", wrap_text=True))
        qty = b["qty"]
        fmt = "#,##0" if isinstance(qty, (int, float)) and float(qty).is_integer() else "#,##0.00"
        _tulis(ws, r, _QTY, qty, font, Alignment(horizontal="right", vertical="top"), fmt)
        _tulis(ws, r, _SAT, b["satuan"], font, Alignment(horizontal="center", vertical="top"))
        _tulis(ws, r, _HARGA, None, font, Alignment(vertical="top"), _AKUNTANSI)
        # Rumus Jumlah di setiap baris ber-QTY, termasuk judul seksi yang ber-QTY ("I DOCKING/
        # GENERAL SERVICE 1 ls" di berkas Basarnas). Baris tanpa QTY tidak diberi rumus: harga
        # yang diisi di sana terlihat tanpa Jumlah, bukan diam-diam dihitung nol.
        jumlah = f"=I{r}*G{r}" if qty is not None else None
        _tulis(ws, r, _JUMLAH, jumlah, font, Alignment(vertical="top"), _AKUNTANSI)
        _tulis(ws, r, _KET, b["keterangan"], Font(name="Arial", size=9, italic=True),
               Alignment(vertical="top", wrap_text=True))
        ws.row_dimensions[r].height = _tinggi(b["uraian"], 2 + t)
    akhir = r0 + len(baris) - 1
    _grid(ws, r0, akhir)
    # Tanpa garis tegak di antara kolom nomor B-F: satu blok uraian, seperti daftar bernomor
    # di Word, bukan kotak-kotak kecil di depan tiap teks.
    for r in range(r0, akhir + 1):
        for c in range(2, _TEKS + 1):
            b = ws.cell(r, c).border
            ws.cell(r, c).border = Border(left=b.left if c == 2 else None,
                                          right=b.right if c == _TEKS else None,
                                          top=b.top, bottom=b.bottom)
    return akhir


def _terbilang(wb, sel_total: str):
    """Sheet tersembunyi berisi rumus terbilang. Sel B6-nya yang dirujuk halaman utama."""
    ws = wb.create_sheet("terbilang")
    ws.sheet_state = "hidden"
    for n in range(1000):
        ws.cell(n + 1, 1, kata_ratusan(n) or None)
    k = "$A$1:$A$1000"
    ws["B1"] = f"=ROUND('Biaya Docking'!{sel_total},0)"
    ws["B2"] = "=INT(B1/1E9)"
    ws["B3"] = "=INT(B1/1E6)-INT(B1/1E9)*1000"
    ws["B4"] = "=INT(B1/1E3)-INT(B1/1E6)*1000"
    ws["B5"] = "=B1-INT(B1/1E3)*1000"
    # ponytail: sampai 999 milyar; di atas itu INDEX keluar rentang dan selnya #REF!.
    ws["B6"] = (
        f'=IF(B1<=0,"",TRIM('
        f'IF(B2>0,INDEX({k},B2+1)&" Milyar ","")'
        f'&IF(B3>0,INDEX({k},B3+1)&" Juta ","")'
        f'&IF(B4=1,"Seribu ",IF(B4>0,INDEX({k},B4+1)&" Ribu ",""))'
        f'&IF(B5>0,INDEX({k},B5+1),""))&" Rupiah")'
    )


def buat(lingkup: dict, tanggal: date) -> bytes:
    """Lingkup dari `repair_list_parser.baca_lingkup()` -> isi berkas .xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Biaya Docking"
    for kol, lebar in _LEBAR.items():
        ws.column_dimensions[kol].width = lebar
    kapal = lingkup["kapal"]

    r0 = _kepala(ws, kapal, tanggal)
    akhir = _isi_tabel(ws, lingkup["baris"], r0)

    # Tepat di bawah item terakhir: Jumlah, PPN 11%, TOTAL. Tidak ada apa pun di antaranya.
    rj = akhir + 1
    for k, (label, rumus) in enumerate((
        ("Jumlah Rp.", f"=SUM(J{r0}:J{akhir})"),
        ("PPN 11% Rp.", f"=J{rj}*11%"),
        ("TOTAL Rp.", f"=J{rj}+J{rj + 1}"),
    )):
        _tulis(ws, rj + k, _HARGA, label, _F_TABEL_TEBAL, Alignment(horizontal="right"))
        _tulis(ws, rj + k, _JUMLAH, rumus, _F_TABEL_TEBAL, fmt=_AKUNTANSI)
    _grid(ws, rj, rj + 2, _HARGA, _JUMLAH)
    ws.cell(rj + 2, _JUMLAH).border = Border(left=_TIPIS, right=_TIPIS, top=_TIPIS, bottom=_GANDA)
    _terbilang(wb, f"J{rj + 2}")

    rt = rj + 4
    ws.merge_cells(start_row=rt, start_column=1, end_row=rt, end_column=5)
    _tulis(ws, rt, 1, "Terbilang :", _F_PITA, Alignment(vertical="center"))
    ws.merge_cells(start_row=rt, start_column=_TEKS, end_row=rt, end_column=_KOLOM_AKHIR)
    _tulis(ws, rt, _TEKS, "=terbilang!$B$6", Font(name="Calibri", size=12, bold=True, italic=True), _TENGAH)
    ws.row_dimensions[rt].height = 32
    _grid(ws, rt, rt, _TEKS, _KOLOM_AKHIR)

    rs = rt + 2
    pemilik, _ = _nilai(kapal, "pemilik")
    _tulis(ws, rs, 2, "Disetujui,")
    _tulis(ws, rs, _HARGA, "Jakarta,", align=Alignment(horizontal="right"))
    _tulis(ws, rs, _JUMLAH, tanggal, fmt="[$-421]dd mmmm yyyy;@", align=Alignment(horizontal="left"))
    _tulis(ws, rs + 1, 2, pemilik, _F_PITA)
    _tulis(ws, rs + 1, _HARGA, "PT. DUKUH RAYA Shipyard", _F_PITA)
    _tulis(ws, rs + 6, _HARGA, PENANDA_TANGAN[0], Font(name="Berlin Sans FB Demi", size=11, underline="single"))
    _tulis(ws, rs + 7, _HARGA, PENANDA_TANGAN[1], Font(name="Calibri", size=11, italic=True))
    ra = rs + 8
    _bingkai(ws, 1, ra)

    ws.sheet_view.showGridLines = False
    ws.print_area = f"A1:K{ra}"
    ws.print_title_rows = "28:29"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
