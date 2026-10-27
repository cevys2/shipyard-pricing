"""File Biaya Docking dari repair list: semua baris ikut, urutannya persis, rumusnya hidup.

Berkas contohnya dibangun di memori, meniru bentuk dan jebakan repair list KMP Gerbang
Samudra 3 2026 -- bukan seluruh isinya. Terhadap berkas aslinya sudah dicek di luar tes
(23 Oktober 2026): 74/74 baris sama dengan Biaya Docking GS3 final buatan tangan, dan setelah
harga finalnya diisi, Excel menghitung TOTAL Rp 733.961.970 -- persis sama -- dengan
terbilang yang benar.
"""

import io
from datetime import date, datetime

import openpyxl
import pytest
from openpyxl import Workbook

from app.services import biaya_docking
from app.services.repair_list_parser import baca_lingkup


def _berkas(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_GS3 = [
    ("REPAIR LIST",),
    ("(Tahun 2026)",),
    (1, "NAMA KAPAL", "KMP. CONTOH"),
    (13, "NAIK DOCK", datetime(2026, 11, 3)),
    ("NO.", "URAIAN", "QTY", "SAT.", "KETERANGAN"),
    ("I", "PELAYANAN UMUM KAPAL"),
    (1, "Docking report (6 Rangkap)", 1, "Ls"),
    ("II", "DOCKING dan UNDOCKING"),
    (4, "Kapal melaksanakan docking dan floating repair"),
    (4.1, "Docking", 14, "Hari"),
    (None, "Pembersihan dan pengecatan"),  # sub-judul tanpa nomor
    # "C" angka romawi yang sah, tapi sesudah seksi II dia bernilai 100, bukan 3.
    ("C", "Main Deck"),
    ("a.", "Valve dilepas dari kedudukannya"),
    (1, "Valve Uk. 10K 200", 1, "buah", "Owner supply"),
    (None, None, None, "Surabaya, 25 September 2026"),
    (None, "Mengetahui,"),  # blok tanda tangan yang mendarat di kolom uraian
]


def test_semua_baris_ikut_persis_urutannya():
    lingkup = baca_lingkup(_berkas(_GS3), "gs3.xlsx")
    assert [(b["no"], b["uraian"], b["qty"], b["satuan"]) for b in lingkup["baris"]] == [
        ("I", "PELAYANAN UMUM KAPAL", None, None),
        ("1", "Docking report (6 Rangkap)", 1, "Ls"),
        ("II", "DOCKING dan UNDOCKING", None, None),
        ("4", "Kapal melaksanakan docking dan floating repair", None, None),
        ("4.1", "Docking", 14, "Hari"),
        ("", "Pembersihan dan pengecatan", None, None),
        ("C", "Main Deck", None, None),
        ("a.", "Valve dilepas dari kedudukannya", None, None),
        ("1", "Valve Uk. 10K 200", 1, "buah"),
    ]
    assert [b["no"] for b in lingkup["baris"] if b["seksi"]] == ["I", "II"]
    assert lingkup["baris"][-1]["keterangan"] == "Owner supply"


def test_ekor_tanda_tangan_dibuang_dan_disebut():
    lingkup = baca_lingkup(_berkas(_GS3), "gs3.xlsx")
    assert all("Mengetahui" not in b["uraian"] for b in lingkup["baris"])
    assert any("Mengetahui" in p for p in lingkup["peringatan"])


def test_pekerjaan_tanpa_qty_di_akhir_tetap_ikut():
    """Laksmana 241: seksi terakhir tanpa QTY sama sekali. Yang dibuang cuma tanda tangan."""
    rows = [
        ("NO", "URAIAN", "QTY", "SAT"),
        ("A", "LAMBUNG KAPAL"),
        (1, "Pengecatan lambung", 224, "m2"),
        ("G", "Fumigasi"),  # bernomor di kolom NO, tanpa QTY
        (None, "3. NDT Vacuum Test Pekerjaan Lambung"),  # nomor di awal teks (Gilimanuk II)
        (None, "Mengetahui,"),
        (None, "PRASETYO ADHI KURNIAWAN,S.T."),
    ]
    lingkup = baca_lingkup(_berkas(rows), "laksmana.xlsx")
    assert [b["uraian"] for b in lingkup["baris"]][-2:] == ["Fumigasi", "3. NDT Vacuum Test Pekerjaan Lambung"]
    assert any("2 baris" in p and "Mengetahui" in p for p in lingkup["peringatan"])


def test_gaya_basarnas_volume_dua_kolom_dan_uraian_melebar():
    """VOL | SAT di baris sub-judul, "SATUAN" di baris judul = kolom HARGA satuan."""
    rows = [
        ("REPAIR LIST",),
        ("NO", "URAIAN PEKERJAAN", None, "VOLUME", None, "SATUAN", "TOTAL", "KET"),
        (None, None, None, "VOL", "SAT"),
        (1, 2, None, 3, 4, 5, 7, 8),  # baris penomoran kolom, bukan data
        ("A", "Pekerjaan Docking/General Service", None, 1, "Ls", 100),
        ("B", "LAMBUNG KAPAL"),
        (None, "a", "Pembersihan lambung", 273, "m²", 20),
        (None, "-", "1 x Primer", 273, "m²", 200),
    ]
    lingkup = baca_lingkup(_berkas(rows), "basarnas.xlsx")
    assert [(b["no"], b["uraian"], b["qty"], b["satuan"]) for b in lingkup["baris"]] == [
        ("A", "Pekerjaan Docking/General Service", 1, "Ls"),
        ("B", "LAMBUNG KAPAL", None, None),
        ("a", "Pembersihan lambung", 273, "m²"),
        ("-", "1 x Primer", 273, "m²"),
    ]


def test_tingkat_gaya_gs3_dibaca_dari_urutan_nomor():
    """Semua nomor di kolom NO. "1" sesudah "b." itu anak b., bukan butir teratas baru."""
    rows = [
        ("NO.", "URAIAN", "QTY", "SAT."),
        ("I", "SEA CHEST dan SEA VALVE"),
        (1, "Kotak Sea Chest"),
        (1.1, "Buka plat sea greating", 5, "buah"),
        (2, "Valve dan strainer"),
        (2.1, "Valve"),
        ("a.", "Valve dilepas dari kedudukannya"),
        ("b.", "Valve sea chest kanan"),
        (1, "Valve Uk. 10K 200", 1, "buah"),
        (2, "Valve Uk. 10K 150", 1, "buah"),
        ("c.", "Valve sea chest kiri"),
        (1, "Valve Uk. 10K 200", 1, "buah"),
        ("II", "LAIN - LAIN"),
        (None, "Pekerjaan tambahan"),  # tanpa nomor, langsung di bawah seksi
        (1, "Cleaning kamar mesin", 10, "m2"),
    ]
    lingkup = baca_lingkup(_berkas(rows), "gs3.xlsx")
    assert [(b["no"], b["tingkat"]) for b in lingkup["baris"]] == [
        ("I", 0), ("1", 0), ("1.1", 1), ("2", 0), ("2.1", 1), ("a.", 2), ("b.", 2),
        ("1", 3), ("2", 3), ("c.", 2), ("1", 3), ("II", 0), ("", 0), ("1", 0),
    ]


def test_tingkat_gaya_basarnas_dibaca_dari_kolomnya():
    """Nomor anak duduk di selnya sendiri di dalam wilayah uraian; kolomnya = tingkatnya."""
    rows = [
        ("NO", "URAIAN", None, None, "QTY", "SAT"),
        (1, "Lambung Kapal"),
        (None, "c", "Perawatan Kotak Laut"),
        (None, None, "1).", "Pembersihan 8 unit lubang", 8, "Unit"),
        (None, None, None, "Material plat aluminium"),  # teks tanpa nomor, sejajar teks 1).
        (None, "d", "Penggantian anode", None, 37, "Unit"),
    ]
    lingkup = baca_lingkup(_berkas(rows), "basarnas.xlsx")
    assert [(b["no"], b["uraian"], b["tingkat"]) for b in lingkup["baris"]] == [
        ("1", "Lambung Kapal", 0),
        ("c", "Perawatan Kotak Laut", 1),
        ("1).", "Pembersihan 8 unit lubang", 2),
        ("", "Material plat aluminium", 2),
        ("d", "Penggantian anode", 1),
    ]


def test_qty_bukan_angka_ditulis_apa_adanya_dan_disebut():
    rows = [("NO", "URAIAN", "QTY", "SAT"), (1, "Cleaning tangki", "1 set", None), (2, "Docking", 3, "Hari")]
    lingkup = baca_lingkup(_berkas(rows), "x.xlsx")
    assert lingkup["baris"][0]["qty"] == "1 set"
    assert any("1 set" in p for p in lingkup["peringatan"])


def test_tanpa_baris_judul_gagal_keras():
    with pytest.raises(ValueError, match="URAIAN"):
        baca_lingkup(_berkas([("Daftar pekerjaan",), ("Cleaning", 1)]), "x.xlsx")


def test_file_biaya_docking_harga_kosong_rumus_hidup():
    lingkup = baca_lingkup(_berkas(_GS3), "gs3.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(biaya_docking.buat(lingkup, date(2026, 10, 23))))
    ws = wb["Biaya Docking"]
    r0 = 30
    akhir = r0 + len(lingkup["baris"]) - 1
    assert ws.cell(r0, 2).value == "PELAYANAN UMUM KAPAL"

    for k, b in enumerate(lingkup["baris"]):
        r = r0 + k
        # Nomor di kolom tingkatnya (A = tingkat 0), teks satu kolom di kanannya.
        assert ws.cell(r, 1 + b["tingkat"]).value == (b["no"] or None)
        assert ws.cell(r, 2 + b["tingkat"]).value == b["uraian"]
        assert ws.cell(r, 9).value is None  # harga satuan: diisi Lutfi di Excel
        assert ws.cell(r, 10).value == (f"=I{r}*G{r}" if b["qty"] is not None else None)
    # "1 Valve Uk. 10K 200" anak "a." anak "C": nomornya di E, teksnya di F.
    assert (ws.cell(akhir, 5).value, ws.cell(akhir, 6).value) == ("1", "Valve Uk. 10K 200")

    # Tepat di bawah item terakhir, tanpa apa pun di antaranya.
    assert [ws.cell(akhir + k, 9).value for k in (1, 2, 3)] == ["Jumlah Rp.", "PPN 11% Rp.", "TOTAL Rp."]
    assert ws.cell(akhir + 1, 10).value == f"=SUM(J{r0}:J{akhir})"
    assert ws.cell(akhir + 2, 10).value == f"=J{akhir + 1}*11%"
    assert ws.cell(akhir + 3, 10).value == f"=J{akhir + 1}+J{akhir + 2}"

    # Terbilang ikut TOTAL lewat sheet tersembunyi, bukan teks mati.
    tb = wb["terbilang"]
    assert tb.sheet_state == "hidden"
    assert tb["B1"].value == f"=ROUND('Biaya Docking'!J{akhir + 3},0)"
    assert ws.cell(akhir + 5, 6).value == "=terbilang!$B$6"
    assert tb.cell(115 + 1, 1).value == "Seratus Lima Belas"


@pytest.mark.parametrize(
    "n, kata",
    [
        (0, ""), (1, "Satu"), (10, "Sepuluh"), (11, "Sebelas"), (12, "Dua Belas"),
        (19, "Sembilan Belas"), (20, "Dua Puluh"), (21, "Dua Puluh Satu"), (100, "Seratus"),
        (101, "Seratus Satu"), (111, "Seratus Sebelas"), (200, "Dua Ratus"),
        (961, "Sembilan Ratus Enam Puluh Satu"), (999, "Sembilan Ratus Sembilan Puluh Sembilan"),
    ],
)
def test_kata_ratusan(n, kata):
    assert biaya_docking.kata_ratusan(n) == kata
