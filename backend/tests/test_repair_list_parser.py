"""Parser REPAIR LIST: lima jebakan nyata + palang rekonsiliasi.

Berkas contohnya dibangun di memori, bukan dibaca dari tiga berkas Basarnas yang asli --
berkas itu ada di luar repo dan tes yang bergantung padanya akan hijau di satu laptop lalu
merah di mana pun yang lain. Yang ditiru di sini bentuk dan jebakannya, satu per satu,
dengan angka yang cukup kecil untuk dijumlahkan dengan mata.

Yang paling penting dijaga di berkas ini adalah **rekonsiliasi**: jumlah semua baris yang
diambil parser harus persis sama dengan angka JUMLAH di dokumen. Semua aturan pengambilan
baris yang lain -- lewati induk yang tidak berharga, lewati rincian "- ..." tanpa harga,
lewati baris penomoran kolom -- tidak bisa dibuktikan benar satu per satu; yang
membuktikannya justru angka totalnya. Kalau satu aturan salah, selisihnya bukan nol.
"""

import io

import pytest
from openpyxl import Workbook

from app.services.repair_list_parser import (
    _ANAK_ANGKA,
    nilai_romawi,
    parse_repair_list_file,
)

# Baris isi: (NO, URAIAN, VOL, SAT, HARGA, TOTAL, KETERANGAN).
# Angkanya sengaja kecil supaya JUMLAH-nya bisa dihitung tangan: 1.232.
_ISI = [
    ("I", "DOCKING/ GENERAL SERVICE", 1, "ls", 100, 100, None),
    (None, None, None, None, None, None, None),
    ("II", "LAMBUNG KAPAL", None, None, None, None, None),
    ("A", "Lambung Kapal Dibawah Garis Air", None, None, None, None, None),
    (1, "Pembersihan lambung", 10, "m2", 15, 150, None),
    # Induk tanpa harga; yang berharga anak-anaknya (jebakan #4).
    (2, "Pengecatan lambung kapal di bawah garis air (incl. Material)", None, None, None, None, None),
    (None, "a. 1 x Cat Primer", 10, "m2", 20, 200, None),
    (None, "b. 2 x Anti Corrosion (AC)", 10, "m2", 20, 200, None),
    # Induk BERHARGA; anaknya cuma rincian isi lump sum dan tidak boleh ikut terambil.
    (3, "Cleaning Tangki", 1, "ls", 30, 30, None),
    (None, "- Cleaning tanki air tawar (6,5 ton)", None, None, None, None, None),
    (None, "- Buka man hole 6 titik", None, None, None, None, None),
    # Jebakan #3: "C" dan "D" adalah huruf sub-seksi, tapi juga angka romawi yang sah.
    ("C", "Main Deck", None, None, None, None, None),
    (1, "Pengecatan lantai main deck", 5, "m2", 20, 100, None),
    ("D", "Top Deck", None, None, None, None, None),
    (1, "Pengecatan lantai top deck", 5, "m2", 20, 100, None),
    ("III", "PERMESINAN", None, None, None, None, None),
    (1, "Perawatan Mesin Induk", None, None, None, None, None),
    (None, "a. Jasa Perawatan Mesin Induk", 3, "unit", 100, 300, None),
    (None, "b. Penggantian Part Mesin Induk", None, None, None, None, None),
    # Jebakan #5: satu berkas menulis "1)", berkas lain "2  " tanpa kurung tutup.
    (None, "1) Penggantian Part Main Engine Tengah", None, None, None, None, None),
    (None, "- Excentric P/N 51.06501-0339", 1, "Pcs", 12, 12, None),
    (None, "2  Penggantian Part M/E Kiri", None, None, None, None, None),
    (None, "- Excentric P/N 51.06501-0339", 1, "Pcs", 15, 15, None),
    (2, "Fumigasi", 1, "ls", 25, 25, "Dilaksanakan oleh Kantor Kesehatan Pelabuhan"),
]

JUMLAH = 1232
PPN = 135.52
TOTAL = 1367.52


def _berkas(isi=None, *, jumlah=JUMLAH, geser=0, judul_ket="KETERANGAN") -> bytes:
    """Repair list di memori.

    `geser` melebarkan grup VOLUME sebanyak N kolom, meniru WIDURA 225 yang menaruh harga
    di kolom F sementara dua berkas lain menaruhnya di kolom E (jebakan #1).
    """
    isi = _ISI if isi is None else isi
    wb = Workbook()
    ws = wb.active
    ws.title = "KN SAR TES 001"
    ws.append(["RINCIAN"])
    ws.append(["PEKERJAAN SPECIAL INSPECTION UNTUK PEMELIHARAAN KN SAR TES"])
    ws.append(["KANTOR PENCARIAN DAN PERTOLONGAN TES"])
    kosong = [None] * geser
    ws.append(["NO", "URAIAN PEKERJAAN", "VOLUME", None, *kosong, "SATUAN", "TOTAL", judul_ket])
    ws.append([None, None, "VOL", "SAT", *kosong, None, None, None])
    # Jebakan #2: baris penomoran kolom, persis di bawah baris sub-judul.
    ws.append([1, 2, 3, 4, *kosong, 5, 6, 7])
    for no, uraian, vol, sat, harga, total, ket in isi:
        ws.append([no, uraian, vol, sat, *kosong, harga, total, ket])
    ws.append([None] * (7 + geser))
    ws.append(["JUMLAH", None, None, None, *kosong, None, jumlah, None])
    ws.append(["PPN 11%", None, None, None, *kosong, None, PPN, None])
    ws.append(["TOTAL", None, None, None, *kosong, None, TOTAL, None])
    ws.append(["Terbilang : ......"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def hasil():
    return parse_repair_list_file(_berkas(), "repair-list.xlsx")


def _cari(hasil, potongan):
    return [i for i in hasil["items"] if potongan in i["uraian"]]


# --- palang rekonsiliasi ---------------------------------------------------------------

def test_jumlah_baris_terbaca_sama_persis_dengan_jumlah_di_dokumen(hasil):
    """Inti seluruh berkas ini. Kalau selisihnya bukan nol, ada baris yang jatuh."""
    r = hasil["rekonsiliasi"]
    assert r["jumlah_dokumen"] == JUMLAH
    assert r["jumlah_terbaca"] == JUMLAH, f"terbaca {r['jumlah_terbaca']}, dokumen {JUMLAH}"
    assert r["selisih"] == 0
    assert r["cocok"] is True


def test_ppn_dan_total_dokumen_ikut_terbaca(hasil):
    assert hasil["rekonsiliasi"]["ppn_dokumen"] == PPN
    assert hasil["rekonsiliasi"]["total_dokumen"] == TOTAL


def test_berkas_yang_bersih_tidak_menghasilkan_peringatan(hasil):
    assert hasil["warnings"] == []


def test_selisih_bukan_nol_ikut_jadi_peringatan():
    """Kalau JUMLAH di dokumen tidak cocok, itu HARUS berbunyi -- bukan lewat diam-diam."""
    h = parse_repair_list_file(_berkas(jumlah=JUMLAH + 500), "x.xlsx")

    assert h["rekonsiliasi"]["cocok"] is False
    assert h["rekonsiliasi"]["selisih"] == -500
    assert any("TIDAK SAMA" in w for w in h["warnings"])


def test_baris_yang_hilang_ketahuan_dari_selisihnya():
    """Simulasi satu baris gagal terbaca: JUMLAH tetap, isinya berkurang."""
    kurang = [b for b in _ISI if b[1] != "Fumigasi"]
    h = parse_repair_list_file(_berkas(kurang), "x.xlsx")

    assert h["rekonsiliasi"]["selisih"] == -25
    assert h["rekonsiliasi"]["cocok"] is False


def test_jumlah_tidak_ketemu_pun_dikatakan():
    wb = Workbook()
    ws = wb.active
    ws.append(["NO", "URAIAN PEKERJAAN", "VOLUME", None, "SATUAN", "TOTAL"])
    ws.append([None, None, "VOL", "SAT", None, None])
    ws.append([1, "Pekerjaan tanpa penutup", 1, "ls", 10, 10])
    buf = io.BytesIO()
    wb.save(buf)

    h = parse_repair_list_file(buf.getvalue(), "x.xlsx")

    assert h["rekonsiliasi"]["jumlah_dokumen"] is None
    assert h["rekonsiliasi"]["cocok"] is False
    assert any("JUMLAH tidak ketemu" in w for w in h["warnings"])


# --- jebakan #1: posisi kolom bergeser --------------------------------------------------

def test_posisi_kolom_harga_yang_bergeser_tetap_terbaca():
    """WIDURA 225 menaruh harga di kolom F, dua berkas lain di kolom E."""
    normal = parse_repair_list_file(_berkas(), "e.xlsx")
    geser = parse_repair_list_file(_berkas(geser=1), "f.xlsx")

    bersih = lambda h: [  # noqa: E731
        (i["kategori"], i["induk_uraian"], i["uraian"], i["volume"], i["satuan"], i["harga"])
        for i in h["items"]
    ]
    assert bersih(geser) == bersih(normal)
    assert geser["rekonsiliasi"]["cocok"] is True


def test_judul_kolom_keterangan_boleh_disingkat_ket():
    """Dua berkas menulis "KETERANGAN", satu menulis "KET"."""
    h = parse_repair_list_file(_berkas(judul_ket="KET"), "x.xlsx")

    (fumigasi,) = _cari(h, "Fumigasi")
    assert fumigasi["keterangan"].startswith("Dilaksanakan oleh")


# --- jebakan #2: baris penomoran kolom --------------------------------------------------

def test_baris_penomoran_kolom_tidak_ikut_jadi_data(hasil):
    """Baris 1|2|3|4|5|6|7 akan masuk sebagai pekerjaan "2" berharga 5, bertotal 6."""
    assert [i for i in hasil["items"] if i["uraian"] == "2"] == []
    assert [i for i in hasil["items"] if i["harga"] == 5] == []


# --- jebakan #3: huruf sub-seksi yang juga angka romawi ---------------------------------

@pytest.mark.parametrize(
    "sub_seksi,uraian",
    [("Main Deck", "Pengecatan lantai main deck"), ("Top Deck", "Pengecatan lantai top deck")],
)
def test_sub_seksi_c_dan_d_tidak_dibaca_sebagai_seksi_baru(hasil, sub_seksi, uraian):
    """"C" bernilai 100 dan "D" bernilai 500 sebagai angka romawi.

    Dibaca apa adanya, keduanya jadi seksi baru dan seluruh baris di bawahnya kehilangan
    kategori "LAMBUNG KAPAL" -- tanpa satu pun error. Pemutusnya urutan: seksi berikutnya
    setelah II adalah III, bukan C.
    """
    (baris,) = _cari(hasil, uraian)

    assert baris["kategori"] == "LAMBUNG KAPAL", "sub-seksi terbaca sebagai seksi baru"
    assert baris["induk_uraian"] == sub_seksi


def test_seksi_romawi_yang_sah_tetap_jadi_kategori(hasil):
    kategori = {i["kategori"] for i in hasil["items"]}
    assert kategori == {"DOCKING/ GENERAL SERVICE", "LAMBUNG KAPAL", "PERMESINAN"}


def test_seksi_yang_ikut_berharga_tetap_terambil(hasil):
    """"I DOCKING/ GENERAL SERVICE 1 ls 100" ada di ketiga berkas dan masuk hitungan."""
    (baris,) = _cari(hasil, "DOCKING/ GENERAL SERVICE")

    assert baris["harga"] == 100
    assert baris["induk_uraian"] is None, "seksi tidak punya induk -- dia sendiri kategorinya"


@pytest.mark.parametrize(
    "token,nilai", [("I", 1), ("II", 2), ("IV", 4), ("V", 5), ("VIII", 8), ("C", 100), ("D", 500)]
)
def test_nilai_romawi(token, nilai):
    assert nilai_romawi(token) == nilai


# --- jebakan #4: induk tanpa harga, anaknya yang berharga -------------------------------

def test_induk_tanpa_harga_tidak_ikut_terambil(hasil):
    assert _cari(hasil, "Pengecatan lambung kapal di bawah garis air") == []


def test_anak_berharga_terambil_dengan_konteks_induknya(hasil):
    (primer,) = _cari(hasil, "1 x Cat Primer")

    assert primer["harga"] == 200 / 10
    assert primer["induk_uraian"] == (
        "Lambung Kapal Dibawah Garis Air › "
        "Pengecatan lambung kapal di bawah garis air (incl. Material)"
    )


def test_rincian_lump_sum_tanpa_harga_dilewati(hasil):
    """"- Cleaning tanki air tawar" adalah isi lump sum induknya, bukan pekerjaan lain."""
    assert _cari(hasil, "Cleaning tanki air tawar") == []
    (induk,) = _cari(hasil, "Cleaning Tangki")
    assert induk["harga"] == 30


# --- jebakan #5: penanda anak yang tidak konsisten --------------------------------------

def test_dua_gaya_penomoran_blok_sama_sama_terbaca(hasil):
    """"1) Penggantian Part..." dan "2  Penggantian Part..." harus setara."""
    excentric = _cari(hasil, "Excentric")

    assert len(excentric) == 2
    assert [i["induk_uraian"].split(" › ")[-1] for i in excentric] == [
        "Penggantian Part Main Engine Tengah",
        "Penggantian Part M/E Kiri",
    ]


def test_uraian_kembar_dibedakan_oleh_induknya_bukan_oleh_harga(hasil):
    """Inti §3 usulan: tiga baris beruraian identik, yang membedakan cuma baris induknya.

    Dan harganya memang beda -- di ANTAREJA 233 Excentric Rp 12,5 juta untuk mesin tengah
    tapi Rp 15 juta untuk kiri dan kanan. Tanpa konteks induk, orang yang membuka katalog
    enam bulan lagi tidak punya cara tahu baris mana untuk mesin mana.
    """
    excentric = _cari(hasil, "Excentric")

    assert {i["uraian"] for i in excentric} == {"Excentric P/N 51.06501-0339"}
    assert sorted(i["harga"] for i in excentric) == [12, 15]
    assert len({i["induk_uraian"] for i in excentric}) == 2


def test_penanda_dibuang_dari_uraian_dan_pindah_ke_induk(hasil):
    """"a." dan "- " cuma nomor urut; yang disimpan pekerjaannya."""
    assert [i["uraian"] for i in _cari(hasil, "Cat Primer")] == ["1 x Cat Primer"]
    assert [i["uraian"] for i in _cari(hasil, "Anti Corrosion")] == ["2 x Anti Corrosion (AC)"]


@pytest.mark.parametrize("teks", ["2 x Anti Fouling (AF)", "3 x Cat Primer", "12 Pcs gasket"])
def test_angka_diikuti_satu_spasi_bukan_penanda_blok(teks):
    """Kalau "2 x Anti Fouling" dianggap penanda, uraiannya terpotong jadi "x Anti Fouling"."""
    assert _ANAK_ANGKA.match(teks) is None


# --- kolom baru + kepala dokumen --------------------------------------------------------

def test_volume_dan_satuan_terbaca_sebagai_angka_dan_teks_terpisah(hasil):
    (baris,) = _cari(hasil, "Pembersihan lambung")

    assert baris["volume"] == 10
    assert baris["satuan"] == "m2"
    assert baris["harga"] == 15
    assert baris["volume"] * baris["harga"] == baris["total_berkas"]


def test_keterangan_ikut_terbaca(hasil):
    (fumigasi,) = _cari(hasil, "Fumigasi")
    assert fumigasi["keterangan"] == "Dilaksanakan oleh Kantor Kesehatan Pelabuhan"


def test_kepala_dokumen_terbaca(hasil):
    assert hasil["detected_jenis_dokumen"] == "RINCIAN"
    assert hasil["detected_nama_perusahaan"] == "KANTOR PENCARIAN DAN PERTOLONGAN TES"
    assert hasil["detected_nama_kapal"] == "KN SAR TES 001", "nama kapal diambil dari nama sheet"


def test_tahun_sengaja_dibiarkan_kosong(hasil):
    """Repair list tidak memuat tahun, dan menebaknya dari nama berkas sudah terbukti salah.

    Yang di nama berkas itu tahun terbit dokumen, bukan tahun pekerjaannya. Lebih baik
    kosong dan diisi manusia daripada terisi angka yang keliru -- angka itu ikut jadi
    prefix ID baris, jadi salahnya permanen.
    """
    assert hasil["detected_tahun"] == ""


def test_berkas_tanpa_baris_judul_ditolak_dengan_jelas():
    wb = Workbook()
    wb.active.append(["ini bukan repair list"])
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(ValueError, match="judul tabel"):
        parse_repair_list_file(buf.getvalue(), "bukan.xlsx")


def test_baris_berharga_tanpa_volume_dihitung_satu_kali_dan_dilaporkan():
    isi = [("I", "Seksi", None, None, None, None, None), (1, "Lump sum tanpa VOL", None, "ls", 40, 40, None)]
    h = parse_repair_list_file(_berkas(isi, jumlah=40), "x.xlsx")

    assert h["rekonsiliasi"]["jumlah_terbaca"] == 40
    assert any("kolom VOL kosong" in w for w in h["warnings"])


def test_vol_kali_harga_beda_dengan_total_di_berkas_ikut_dilaporkan():
    """Berkas asli boleh saja punya sel TOTAL yang tidak konsisten; jangan dipilih diam-diam."""
    isi = [("I", "Seksi", None, None, None, None, None), (1, "Salah hitung", 2, "unit", 10, 999, None)]
    h = parse_repair_list_file(_berkas(isi, jumlah=20), "x.xlsx")

    assert h["rekonsiliasi"]["jumlah_terbaca"] == 20, "yang disimpan VOL x harga"
    assert any("kolom TOTAL di berkas" in w for w in h["warnings"])
